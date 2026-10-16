"""Salida: un JSON por documento (formato de la especificación) y un Excel consolidado."""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ESTADOS = {
    "COMPLETED": "Correcto",
    "COMPLETED_WITH_WARNINGS": "Con observaciones",
    "VALIDATION_FAILED": "No cuadra",
    "NEW_FORMAT_REVIEW": "Formato nuevo: revisar",
    "TEMPLATE_NOT_IDENTIFIED": "Formato nuevo",
    "ERROR": "Error de lectura",
}
COLOR_ESTADO = {"COMPLETED": "D9EAD3", "COMPLETED_WITH_WARNINGS": "FFF2CC", "VALIDATION_FAILED": "F4CCCC",
                "TEMPLATE_NOT_IDENTIFIED": "D9D2E9", "NEW_FORMAT_REVIEW": "D9D2E9", "ERROR": "F4CCCC"}

COLUMNAS_ORDEN = [
    ("Archivo", "file"), ("Estado", "estado"), ("Plantilla", "plantilla"), ("N° orden", "numero_orden"),
    ("Revisión", "revision"), ("Fecha emisión", "fecha_emision"), ("Agencia", "agencia"), ("RUC agencia", "agencia_ruc"),
    ("Anunciante", "anunciante"), ("RUC anunciante", "anunciante_ruc"), ("Producto", "producto"),
    ("Campaña", "campana"), ("Motivo", "motivo"), ("Medio", "medio"), ("Proveedor", "proveedor"),
    ("RUC proveedor", "proveedor_ruc"), ("Cobertura", "ciudad"), ("Periodo inicio", "periodo_inicio"),
    ("Periodo fin", "periodo_fin"), ("Mes", "mes"), ("Plan", "plan"), ("Presupuesto", "presupuesto"),
    ("Pauta", "pauta"), ("Contrato", "contrato"), ("Condición pago", "condicion_pago"), ("Moneda", "moneda"),
    ("Filas", "filas"), ("Avisos", "avisos"), ("Costo", "costo"), ("Descuento", "descuento"), ("Neto", "neto"),
    ("IGV", "igv"), ("Total", "total"), ("Emitido por", "emitido_por"), ("Errores", "errores"),
    ("Observaciones", "observaciones_val"),
]


def _ancho(ws, maximo=60):
    for col in ws.columns:
        largo = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(10, largo + 2), maximo)


def _cabecera(ws, titulos):
    ws.append(titulos)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def _val(campo):
    return campo.get("value") if isinstance(campo, dict) else campo


def libro(resultados: list[dict]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Órdenes"
    _cabecera(ws, [t for t, _ in COLUMNAS_ORDEN])
    wd = wb.create_sheet("Detalle")
    _cabecera(wd, ["Archivo", "N° orden", "Fila", "Sección", "Paquete", "Descripción", "Motivo", "Horario",
                   "Duración (s)", "Avisos", "Costo unitario", "Costo total", "Bonificado", "Días con avisos"])
    wp = wb.create_sheet("Programación")
    _cabecera(wp, ["Archivo", "N° orden", "Medio", "Fila", "Descripción", "Fecha", "Día", "Avisos"])
    wv = wb.create_sheet("Validaciones")
    _cabecera(wv, ["Archivo", "N° orden", "Estado", "Severidad", "Regla", "Fila", "Qué se revisa", "Detalle"])

    for r in resultados:
        s = r.get("summary") or {}
        rep = r.get("validationReport") or {}
        fallas = [x for x in rep.get("executedRules", []) if x.get("passed") is False]
        fila = dict(s)
        fila.update(file=r["file"], estado=ESTADOS.get(r["status"], r["status"]),
                    plantilla=(r.get("template") or {}).get("templateId"), errores=rep.get("errors"),
                    observaciones_val="; ".join(f"{x['rule']}: {x.get('detail') or x.get('difference')}" for x in fallas))
        ws.append([fila.get(k) for _, k in COLUMNAS_ORDEN])
        color = COLOR_ESTADO.get(r["status"])
        if color:
            ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor=color)

        datos = r.get("extractedData") or {}
        for i, f in enumerate(datos.get("detail", []), 1):
            prog = f.get("programacion", [])
            wd.append([r["file"], s.get("numero_orden"), i, _val(f.get("seccion")), _val(f.get("paquete")),
                       _val(f.get("descripcion")), _val(f.get("motivo")), _val(f.get("horario")),
                       _val(f.get("duracion")), _val(f.get("total_avisos")), _val(f.get("costo_unitario")),
                       _val(f.get("costo_total")), "Sí" if _val(f.get("bonificado")) else "",
                       ", ".join(str(p["dia"]) for p in prog)])
            for p in prog:
                wp.append([r["file"], s.get("numero_orden"), s.get("medio"), i, _val(f.get("descripcion")),
                           p.get("fecha"), p.get("letra"), p["avisos"]])
        for x in rep.get("executedRules", []):
            wv.append([r["file"], s.get("numero_orden"),
                       {"PASSED": "Cumple", "FAILED": "No cumple", "SKIPPED": "No aplica", "ERROR": "Error"}.get(x["status"], x["status"]),
                       x.get("severity"), x["rule"], x.get("row"), x.get("description"),
                       x.get("detail") if x.get("detail") else (x.get("difference") if x["status"] == "FAILED" else None)])
            if x["status"] == "FAILED":
                wv.cell(wv.max_row, 3).fill = PatternFill("solid", fgColor="F4CCCC" if x["severity"] == "ERROR" else "FFF2CC")
    for hoja in (ws, wd, wp, wv):
        _ancho(hoja)
        hoja.auto_filter.ref = hoja.dimensions
    for col in ("AC", "AD", "AE", "AF", "AG"):
        for c in ws[col][1:]:
            c.number_format = "#,##0.00"
    for col in ("K", "L"):
        for c in wd[col][1:]:
            c.number_format = "#,##0.00"
    return wb


def guardar_resultados(resultados: list[dict], carpeta: Path) -> dict:
    carpeta = Path(carpeta)
    carpeta.mkdir(parents=True, exist_ok=True)
    for r in resultados:
        limpio = {k: v for k, v in r.items() if not k.startswith("_")}
        nombre = Path(r["file"]).stem + ".json"
        (carpeta / nombre).write_text(json.dumps(limpio, ensure_ascii=False, indent=2), encoding="utf-8")
    excel = carpeta / f"ordenes_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    libro(resultados).save(excel)
    return {"carpeta": str(carpeta), "excel": str(excel)}
